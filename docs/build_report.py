from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).parent; E=ROOT/'report-evidence'; SUPPORT=ROOT/'report-support'
E.mkdir(exist_ok=True)
md=ROOT/'report-chapters-draft.md'
# Keep the public build independent of private citation tools and full-text archives.
# This gate checks citation identity and URLs, not the truth of every cited claim.
ledger=json.loads((SUPPORT/'citation-ledger.json').read_text())
source_text=md.read_text().split('## Sources',1)
if len(source_text)!=2:raise ValueError('The report must contain one Sources section')
refs={int(n):url for n,url in re.findall(r'^\[(\d+)\] (https?://\S+) — ',source_text[1],re.M)}
expected={source['id']:source['url'] for source in ledger['sources']}
if refs!=expected:raise ValueError('Report references differ from the citation ledger')
used={int(n) for n in re.findall(r'\[(\d+)\]',source_text[0])}
if not used.issubset(refs):raise ValueError('Report contains an unresolved citation number')
(E/'citation-verification.txt').write_text('Reference numbers and URLs match the archived citation ledger.\nThis is a consistency check, not fresh full-text source verification.\n')
d=Document();sec=d.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=sec.bottom_margin=Inches(1.5);sec.left_margin=sec.right_margin=Inches(1)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12 if name=='Normal' else 16 if name in ['Title','Heading 1'] else 14)
 st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.5
 if name!='Normal':st.font.bold=True
 # Remove the template's decorative Title border, which otherwise renders blue.
 for border in list(st.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
# Word's built-in theme font attributes override explicit fonts in some renderers.
# Apply the university font to paragraph and linked character styles consistently.
for name in ['Normal','Title','Title Char','Heading 1','Heading 1 Char','Heading 2','Heading 2 Char','Heading 3','Heading 3 Char','Caption','Header','Footer']:
 if name not in d.styles:continue
 st=d.styles[name];st.font.name='Times New Roman'
 fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
 for attr in list(fonts.attrib):
  if 'theme' in attr.lower():del fonts.attrib[attr]
 for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
 st.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=1
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
sec.header.paragraphs[0].text='HLL-04 | UI design for data collection'
body_count=0;inbody=False;previous_table=False
paragraphs=[]
for block in md.read_text().split('\n\n'):
 if block.startswith('[1] '): paragraphs.extend(block.splitlines())
 else: paragraphs.append(block)
for para in paragraphs:
 para=para.strip()
 if not para:continue
 if para.startswith('# '):d.add_paragraph(para[2:],style='Title');continue
 if para.startswith('## '):
  title=para[3:];h=d.add_heading(title,1)
  if title=='Abstract' or re.match(r'[1-6] ',title): h.paragraph_format.page_break_before=True
  inbody=bool(re.match(r'[1-6] ',title));continue
 if para.startswith('### '):d.add_heading(para[4:],2);continue
 if re.match(r'^Table \d+\. ',para):
  p=d.add_paragraph(para,style='Caption');p.paragraph_format.keep_with_next=True
  p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(6)
  for rr in p.runs:rr.font.name='Times New Roman';rr.font.size=Pt(11);rr.font.color.rgb=RGBColor(0,0,0)
  continue
 if para.startswith('![Figure '):
  match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',para)
  if not match: raise ValueError('Invalid figure: '+para)
  figure_width=5.5 if match[2]=='report-support/review-details.png' else 6.1
  d.add_picture(str(ROOT/match[2]),width=Inches(figure_width))
  picture_paragraph=d.paragraphs[-1]
  picture_paragraph.paragraph_format.keep_with_next=True
  if match[2]=='report-support/review-details.png':
   picture_paragraph.alignment=1
   picture_paragraph.paragraph_format.line_spacing=1
  cap=d.add_paragraph(match[1],style='Caption');cap.paragraph_format.keep_with_next=False
  for rr in cap.runs:rr.font.size=Pt(11)
  continue
 if para.startswith('|'):
  rows=[[cell.strip() for cell in line.strip().strip('|').split('|')] for line in para.splitlines()]
  rows=[row for row in rows if not all(re.fullmatch(r'[-:]+',cell) for cell in row)]
  table=d.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
  pr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
  pr.append(borders)
  for i,row in enumerate(rows):
   cells=table.add_row().cells
   for cell,text in zip(cells,row):
    cell.text=text
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcp=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'80');margin.set(qn('w:type'),'dxa');margins.append(margin)
    tcp.append(margins)
    if i==0:
     shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF1');tcp.append(shading)
    for pp in cell.paragraphs:
     pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(4)
     for rr in pp.runs:rr.font.size=Pt(10);rr.bold=(i==0)
   trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if i==0:trpr.append(OxmlElement('w:tblHeader'))
  previous_table=True
  continue
 if inbody:body_count+=len(para.split())
 p=d.add_paragraph()
 if previous_table:p.paragraph_format.space_before=Pt(6)
 previous_table=False
 if not inbody and 'title' in locals() and title=='Use of Generative AI Tools — draft disclosure':
  p.paragraph_format.line_spacing=1.15
 ref=re.match(r'^(\[\d+\]) (https?://\S+) — (.*)$',para)
 if ref:
  para=f'{ref[1]} {ref[3]}. [Source text]({ref[2]})'
  p.paragraph_format.line_spacing=1
  p.paragraph_format.keep_together=True
  p.paragraph_format.space_after=Pt(6)
 # Render Markdown source links as live clickable hyperlinks.
 parts=re.split(r'(\[[^\]]+\]\(https?://[^\s]+\))',para)
 for part in parts:
  m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^\s]+)\)',part)
  if not m:p.add_run(part);continue
  rel=d.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color);run.append(props);text=OxmlElement('w:t');text.text=m[1];run.append(text);h.append(run);p._p.append(h)
 if para.startswith('- ['):p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
out=ROOT/'report-chapters-draft.docx';d.save(out)
# Compute body count from final Word XML including ordinary text only, excluding headings and front/end matter.
loaded=Document(out);active=False;count=0
for p in loaded.paragraphs:
 if p.style.name=='Heading 1':active=bool(re.match(r'[1-6] ',p.text))
 elif active and not p.style.name.startswith('Heading') and p.style.name!='Caption':count+=len(p.text.split())
assert count==body_count
(E/'document-statistics.json').write_text(json.dumps({'body_words_whitespace':count,'rule':'Chapters 1–6 prose only; excludes headings, tables, identity, draft notice, AI disclosure, references, header/footer','docx_paragraphs':len(loaded.paragraphs)},indent=2))
print('BODY WORDS',count,'DOCX',out)
